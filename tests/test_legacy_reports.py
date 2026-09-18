import datetime as dt
import shutil
import subprocess
import tempfile
import unittest
import zipfile
from pathlib import Path
from unittest.mock import Mock, patch
from xml.etree import ElementTree as ET

from openpyxl import load_workbook

from monthly_workbook.legacy_reports import (
    ReportKind,
    SupplierChange,
    convert_xls,
    detect_report,
    update_legacy_report,
)


DESKTOP = Path("/Users/tongchengjiushui/Desktop")
SOURCES = {
    ReportKind.CHEMICAL_USAGE: DESKTOP / "开阳磷矿供贵阳化肥矿石用量(9月).xls",
    ReportKind.INVENTORY_PLAN: DESKTOP / "大水矿石产消存计划一览表（9月）.xls",
    ReportKind.DISPATCH_INFO: DESKTOP / "开阳大水工业园区调度生产信息（9月）.xls",
}


class WindowsConversionTests(unittest.TestCase):
    def test_convert_xls_uses_installed_windows_office_when_libreoffice_is_missing(self):
        with tempfile.TemporaryDirectory() as directory:
            source = Path(directory) / "旧报表.xls"
            source.write_bytes(b"xls")
            destination = Path(directory) / "converted"

            def create_output(command, **_kwargs):
                destination.mkdir(parents=True, exist_ok=True)
                (destination / "旧报表.xlsx").write_bytes(b"xlsx")
                return Mock(returncode=0, stdout="", stderr="")

            with (
                patch("monthly_workbook.legacy_reports._libreoffice_path", return_value=None),
                patch("monthly_workbook.legacy_reports._windows_conversion_script", return_value=Path("convert_xls_windows.ps1")),
                patch("monthly_workbook.legacy_reports.sys.platform", "win32"),
                patch("monthly_workbook.legacy_reports.subprocess.run", side_effect=create_output) as run,
            ):
                result = convert_xls(source, destination)

            self.assertEqual(result, destination.resolve() / "旧报表.xlsx")
            command = run.call_args.args[0]
            self.assertEqual(command[0], "powershell.exe")
            self.assertIn("convert_xls_windows.ps1", command)
            self.assertIn(str(source.resolve()), command)

    def test_convert_xls_explains_requirements_when_windows_office_conversion_fails(self):
        with tempfile.TemporaryDirectory() as directory:
            source = Path(directory) / "旧报表.xls"
            source.write_bytes(b"xls")
            with (
                patch("monthly_workbook.legacy_reports._libreoffice_path", return_value=None),
                patch("monthly_workbook.legacy_reports._windows_conversion_script", return_value=Path("convert_xls_windows.ps1")),
                patch("monthly_workbook.legacy_reports.sys.platform", "win32"),
                patch(
                    "monthly_workbook.legacy_reports.subprocess.run",
                    return_value=Mock(returncode=1, stdout="", stderr="未找到可用的表格程序"),
                ),
            ):
                with self.assertRaisesRegex(RuntimeError, "Excel、WPS 或 LibreOffice"):
                    convert_xls(source, Path(directory) / "converted")


class LegacyReportTests(unittest.TestCase):
    @classmethod
    def setUpClass(cls):
        converted_dir = Path("/tmp/monthly-report-inspect")
        if not shutil.which("soffice"):
            raise unittest.SkipTest("缺少 LibreOffice")
        cls.directory = tempfile.TemporaryDirectory()
        cls.converted = {}
        for kind, path in SOURCES.items():
            if path.exists():
                cls.converted[kind] = convert_xls(path, Path(cls.directory.name))
            else:
                fallback = converted_dir / f"{path.stem}.xlsx"
                if not fallback.exists():
                    raise unittest.SkipTest(f"缺少测试模板：{path.name}")
                cls.converted[kind] = fallback

    @classmethod
    def tearDownClass(cls):
        cls.directory.cleanup()

    def test_detects_three_report_types(self):
        for expected, path in self.converted.items():
            with self.subTest(path=path.name):
                self.assertEqual(detect_report(path), expected)

    def test_chemical_usage_updates_dates_references_and_suppliers(self):
        output = Path(self.directory.name) / "chemical-result.xlsx"
        update_legacy_report(
            self.converted[ReportKind.CHEMICAL_USAGE], output,
            dt.date(2026, 9, 25), dt.date(2026, 10, 24),
            [
                SupplierChange("rename", "供应商", "司尔特", "测试供应商"),
                SupplierChange("add", "供应商", "", "新增供应商"),
                SupplierChange("add", "供应商", "", "第二新增供应商"),
                SupplierChange("delete", "供应商", "新增供应商", ""),
            ],
        )
        workbook = load_workbook(output, data_only=False)
        date_sheets = [sheet for sheet in workbook if sheet["F2"].value == "填表日期："]
        self.assertEqual(len(date_sheets), 30)
        self.assertEqual(date_sheets[0].title, "9.25")
        self.assertEqual(date_sheets[0]["G2"].value, "2026.9.25")
        self.assertIsNone(date_sheets[0]["D6"].value)
        self.assertIsNone(date_sheets[0]["C6"].value)
        self.assertEqual(date_sheets[1]["E4"].value, "=D4+'9.25'!E4")
        self.assertIn("测试供应商", [cell.value for cell in date_sheets[0]["B"]])
        self.assertNotIn("新增供应商", [cell.value for cell in date_sheets[0]["B"]])
        self.assertIn("第二新增供应商", [cell.value for cell in date_sheets[0]["B"]])
        final_added_row = next(cell.row for cell in date_sheets[0]["B"] if cell.value == "第二新增供应商")
        self.assertIn(f"A4:A{final_added_row}", {str(value) for value in date_sheets[0].merged_cells.ranges})
        self.assertIsNone(date_sheets[0].cell(final_added_row, 1).value)
        total_row = next(cell.row for cell in date_sheets[0]["A"] if cell.value == "合计")
        self.assertEqual(date_sheets[0].cell(total_row, 4).value, f"=SUM(D4:D{total_row - 1})")

    def test_dispatch_updates_dates_and_adds_supplier(self):
        output = Path(self.directory.name) / "dispatch-result.xlsx"
        update_legacy_report(
            self.converted[ReportKind.DISPATCH_INFO], output,
            dt.date(2026, 9, 25), dt.date(2026, 10, 24),
            [
                SupplierChange("add", "外购供应商", "", "待删除供应商"),
                SupplierChange("add", "外购供应商", "", "新增外购供应商"),
                SupplierChange("delete", "外购供应商", "待删除供应商", ""),
            ],
        )
        workbook = load_workbook(output, data_only=False)
        date_sheets = [sheet for sheet in workbook if sheet["G4"].value == "日期："]
        self.assertEqual(date_sheets[0].title, "9.25")
        self.assertEqual(date_sheets[0]["I4"].value, "2026.9.25")
        self.assertIsNone(date_sheets[0]["B14"].value)
        self.assertIsNone(date_sheets[0]["H14"].value)
        row = next(cell.row for cell in date_sheets[0]["A"] if cell.value == "新增外购供应商")
        self.assertEqual(date_sheets[0].cell(row, 3).value, f"=B{row}")
        self.assertEqual(date_sheets[1].cell(row, 3).value, f"=B{row}+'9.25'!C{row}")
        self.assertEqual(date_sheets[0].row_dimensions[row].height, date_sheets[0].row_dimensions[row - 1].height)
        self.assertEqual(date_sheets[0].cell(row, 12)._style, date_sheets[0].cell(row - 1, 12)._style)
        self.assertIn(f"E14:E{row}", {str(value) for value in date_sheets[0].merged_cells.ranges})
        self.assertEqual(date_sheets[0].cell(row, 5).border.right.style, "thin")
        self.assertEqual(date_sheets[0].cell(row, 5).border.bottom.style, "thin")
        self.assertEqual(date_sheets[0].cell(row, 5).fill.fgColor.rgb, date_sheets[0].cell(row - 1, 5).fill.fgColor.rgb)

    def test_inventory_plan_chooses_supplier_area_without_changing_formulas(self):
        output = Path(self.directory.name) / "inventory-result.xlsx"
        source = self.converted[ReportKind.INVENTORY_PLAN]
        before = load_workbook(source, data_only=False)
        active_before = max(before.worksheets, key=lambda ws: sum(isinstance(ws.cell(2, col).value, dt.datetime) for col in range(6, 37)))
        update_legacy_report(
            source, output, dt.date(2026, 9, 25), dt.date(2026, 10, 24),
            [
                SupplierChange("add", "生产日累", "", "新增生产供应商"),
                SupplierChange("add", "当日配送量", "", "待删除配送商"),
                SupplierChange("add", "当日配送量", "", "新增配送供应商"),
                SupplierChange("delete", "当日配送量", "待删除配送商", ""),
            ],
        )
        workbook = load_workbook(output, data_only=False)
        active = max(
            workbook.worksheets,
            key=lambda ws: max(
                (ws.cell(2, col).value for col in range(6, 37) if isinstance(ws.cell(2, col).value, (dt.date, dt.datetime))),
                default=dt.datetime.min,
            ),
        )
        self.assertEqual(active["F2"].value.date(), dt.date(2026, 9, 25))
        self.assertIsNone(active["F3"].value)
        self.assertIsNone(active["F53"].value)
        note_row = next(cell.row for cell in active["B"] if cell.value == "填表说明：")
        for row in range(3, note_row):
            self.assertEqual(active.cell(row, 37).value, f"=SUM(F{row}:AJ{row})")
        production_row = next(cell.row for cell in active["B"] if cell.value == "新增生产供应商")
        delivery_row = next(cell.row for cell in active["B"] if cell.value == "新增配送供应商")
        merges = {str(value) for value in active.merged_cells.ranges}
        self.assertIn(f"B{production_row}:E{production_row}", merges)
        self.assertIn(f"B{delivery_row}:E{delivery_row}", merges)
        self.assertTrue(any(value.startswith("A3:A") and value.endswith(str(production_row)) for value in merges))
        self.assertTrue(any(value.startswith("AL") and value.endswith(f":AL{delivery_row}") for value in merges))
        self.assertEqual(active.cell(production_row, 37).value, f"=SUM(F{production_row}:AJ{production_row})")
        self.assertEqual(active.cell(delivery_row, 37).value, f"=SUM(F{delivery_row}:AJ{delivery_row})")
        self.assertEqual(
            active.cell(production_row, 38)._style,
            active.cell(production_row - 1, 38)._style,
        )
        self.assertIsNone(active.cell(production_row, 38).value)
        for row in (production_row, delivery_row):
            self.assertEqual(active.cell(row, 2).border.left.style, "thin")
            self.assertEqual(active.cell(row, 2).border.top.style, "thin")
            self.assertEqual(active.cell(row, 2).border.bottom.style, "thin")
            self.assertEqual(active.cell(row, 5).border.right.style, "thin")
            self.assertEqual(active.cell(row, 5).border.top.style, "thin")
            self.assertEqual(active.cell(row, 5).border.bottom.style, "thin")
        self.assertIn("新增配送供应商", [cell.value for cell in active["B"]])
        self.assertNotIn("待删除配送商", [cell.value for cell in active["B"]])
        for sheet in workbook.worksheets:
            for row in sheet.iter_rows():
                for cell in row:
                    if isinstance(cell.value, str) and cell.value.startswith("="):
                        self.assertNotIn("#REF!", cell.value)

        with zipfile.ZipFile(output) as archive:
            workbook_xml = ET.fromstring(archive.read("xl/workbook.xml"))
            relationships = ET.fromstring(archive.read("xl/_rels/workbook.xml.rels"))
            namespace = {"m": "http://schemas.openxmlformats.org/spreadsheetml/2006/main"}
            rel_namespace = {"r": "http://schemas.openxmlformats.org/package/2006/relationships"}
            rel_id = next(
                node.attrib["{http://schemas.openxmlformats.org/officeDocument/2006/relationships}id"]
                for node in workbook_xml.findall("m:sheets/m:sheet", namespace)
                if node.attrib["name"] == active.title
            )
            target = next(node.attrib["Target"] for node in relationships.findall("r:Relationship", rel_namespace) if node.attrib["Id"] == rel_id)
            sheet_xml = ET.fromstring(archive.read(target.lstrip("/")))
            cell = next(node for node in sheet_xml.findall(".//m:c", namespace) if node.attrib.get("r") == f"AK{delivery_row}")
            self.assertEqual(cell.findtext("m:v", namespaces=namespace), "0")


if __name__ == "__main__":
    unittest.main()
