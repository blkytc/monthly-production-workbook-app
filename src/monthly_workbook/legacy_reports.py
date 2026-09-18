"""Update dates and supplier rows in the three legacy monthly reports."""

from __future__ import annotations

import copy
import datetime as dt
import re
import shutil
import subprocess
import zipfile
from dataclasses import dataclass
from enum import Enum
from pathlib import Path

from openpyxl import load_workbook
from openpyxl.formula.translate import Translator
from openpyxl.styles import Border, Side
from openpyxl.utils import get_column_letter, range_boundaries

from .sizing import column_width, parse_columns, parse_rows, row_height
from .workbook import DimensionRule


class ReportKind(str, Enum):
    CHEMICAL_USAGE = "开阳磷矿供贵阳化肥矿石用量"
    INVENTORY_PLAN = "大水矿石产消存计划一览表"
    DISPATCH_INFO = "开阳大水工业园区调度生产信息"


@dataclass(frozen=True)
class SupplierChange:
    action: str
    area: str
    old_name: str
    new_name: str


def _soffice() -> str:
    candidates = [
        shutil.which("soffice"),
        "/Applications/LibreOffice.app/Contents/MacOS/soffice",
        str(Path.home() / ".cache/codex-runtimes/codex-primary-runtime/dependencies/bin/override/soffice"),
        r"C:\Program Files\LibreOffice\program\soffice.exe",
    ]
    for candidate in candidates:
        if candidate and Path(candidate).exists():
            return str(candidate)
    raise RuntimeError("处理 .xls 需要安装 LibreOffice，或先用 Excel/WPS 另存为 .xlsx")


def convert_xls(source: str | Path, output_dir: str | Path) -> Path:
    source_path = Path(source).expanduser().resolve()
    output_path = Path(output_dir).expanduser().resolve()
    output_path.mkdir(parents=True, exist_ok=True)
    profile = output_path / ".libreoffice-profile"
    command = [
        _soffice(), f"-env:UserInstallation={profile.as_uri()}", "--headless",
        "--convert-to", "xlsx", "--outdir", str(output_path), str(source_path),
    ]
    result = subprocess.run(command, capture_output=True, text=True, timeout=120)
    converted = output_path / f"{source_path.stem}.xlsx"
    if result.returncode or not converted.exists():
        detail = result.stderr.strip() or result.stdout.strip() or "转换失败"
        raise RuntimeError(f"无法转换旧版 Excel 文件：{detail}")
    return converted


def detect_report(path: str | Path) -> ReportKind:
    workbook = load_workbook(path, read_only=True, data_only=False)
    try:
        for sheet in workbook.worksheets:
            values = " ".join(str(sheet.cell(row, column).value or "") for row in range(1, 5) for column in range(1, 11))
            if "供贵阳化肥矿石用量" in values:
                return ReportKind.CHEMICAL_USAGE
            if "大水工业园区矿石产消存计划一览表" in values:
                return ReportKind.INVENTORY_PLAN
            if "大水工业园区调度生产信息" in values:
                return ReportKind.DISPATCH_INFO
    finally:
        workbook.close()
    raise ValueError("未识别出受支持的旧版月度报表")


def report_title(path: str | Path, kind: ReportKind | None = None) -> str:
    report = kind or detect_report(path)
    workbook = load_workbook(path, read_only=False, data_only=False)
    try:
        if report == ReportKind.INVENTORY_PLAN:
            return str(_active_inventory_sheet(workbook)["A1"].value or "")
        sheets = _daily_sheets(workbook, report)
        address = "A1" if report == ReportKind.CHEMICAL_USAGE else "A3"
        return str(sheets[0][address].value or "")
    finally:
        workbook.close()


def _dates(start: dt.date, end: dt.date) -> list[dt.date]:
    if end < start:
        raise ValueError("结束日期不能早于开始日期")
    values = [start + dt.timedelta(days=offset) for offset in range((end - start).days + 1)]
    if len(values) > 31:
        raise ValueError("这类月度报表一次最多支持 31 天")
    return values


def _copy_row_style(sheet, source_row: int, target_row: int, max_column: int) -> None:
    source_height = sheet.row_dimensions[source_row].height
    sheet.row_dimensions[target_row].height = source_height
    for column in range(1, max_column + 1):
        source = sheet.cell(source_row, column)
        target = sheet.cell(target_row, column)
        if source.has_style:
            target._style = copy.copy(source._style)
        target.number_format = source.number_format
        target.alignment = copy.copy(source.alignment)
        target.protection = copy.copy(source.protection)


def _insert_row_preserving_merges(sheet, row_number: int) -> None:
    shifted = []
    for merged in list(sheet.merged_cells.ranges):
        min_col, min_row, max_col, max_row = range_boundaries(str(merged))
        if min_row >= row_number:
            sheet.unmerge_cells(str(merged))
            shifted.append((min_col, min_row + 1, max_col, max_row + 1))
    sheet.insert_rows(row_number)
    for min_col, min_row, max_col, max_row in shifted:
        sheet.merge_cells(
            f"{get_column_letter(min_col)}{min_row}:{get_column_letter(max_col)}{max_row}"
        )


def _delete_row_preserving_merges(sheet, row_number: int) -> None:
    adjusted = []
    for merged in list(sheet.merged_cells.ranges):
        min_col, min_row, max_col, max_row = range_boundaries(str(merged))
        sheet.unmerge_cells(str(merged))
        if min_row > row_number:
            min_row -= 1
            max_row -= 1
        elif min_row <= row_number <= max_row:
            max_row -= 1
        if min_row <= max_row:
            adjusted.append((min_col, min_row, max_col, max_row))
    sheet.delete_rows(row_number)
    for row in range(row_number, sheet.max_row + 1):
        for column in range(1, sheet.max_column + 1):
            cell = sheet.cell(row, column)
            if isinstance(cell.value, str) and cell.value.startswith("="):
                old_coordinate = f"{get_column_letter(column)}{row + 1}"
                try:
                    cell.value = Translator(cell.value, origin=old_coordinate).translate_formula(cell.coordinate)
                except Exception:
                    pass
    for min_col, min_row, max_col, max_row in adjusted:
        sheet.merge_cells(
            f"{get_column_letter(min_col)}{min_row}:{get_column_letter(max_col)}{max_row}"
        )


def _extend_section_merge(sheet, column: int, section_row: int, end_row: int) -> None:
    section_merge = next(
        (
            merged for merged in list(sheet.merged_cells.ranges)
            if merged.min_col == column == merged.max_col and merged.min_row == section_row
        ),
        None,
    )
    if section_merge is not None:
        sheet.unmerge_cells(str(section_merge))
    sheet.merge_cells(
        start_row=section_row, start_column=column, end_row=end_row, end_column=column
    )


def _merge_inventory_supplier_row(sheet, row: int, section_row: int) -> None:
    sheet.merge_cells(start_row=row, start_column=2, end_row=row, end_column=5)
    thin = Side(style="thin", color="000000")
    sheet.cell(row, 2).border = Border(left=thin, right=thin, top=thin, bottom=thin)
    _extend_section_merge(sheet, 1, section_row, row)


def _ensure_inventory_monthly_formulas(sheet) -> None:
    note_row = next(cell.row for cell in sheet["B"] if cell.value == "填表说明：")
    for row in range(3, note_row):
        cell = sheet.cell(row, 37)
        if cell.value is None:
            cell.value = f"=SUM(F{row}:AJ{row})"


def _cache_formula_zeros(path: Path, sheet_part: str, addresses: list[str]) -> None:
    if not addresses:
        return
    part = sheet_part.lstrip("/")
    temporary = path.with_suffix(".cached.xlsx")
    with zipfile.ZipFile(path) as source, zipfile.ZipFile(temporary, "w", zipfile.ZIP_DEFLATED) as target:
        for item in source.infolist():
            content = source.read(item.filename)
            if item.filename == part:
                for address in addresses:
                    pattern = rb'(<c\b[^>]*\br="' + re.escape(address.encode()) + rb'"[^>]*>.*?<v>)(?:</v>)'
                    content, count = re.subn(pattern, rb'\g<1>0</v>', content, count=1, flags=re.DOTALL)
                    if count != 1:
                        raise RuntimeError(f"无法写入月累初始值：{address}")
            target.writestr(item, content)
    temporary.replace(path)


def _replace_sheet_references(formula: str, previous_title: str) -> str:
    return re.sub(r"'(?:[^']|'')+'!", f"'{previous_title.replace("'", "''")}'!", formula)


def _daily_sheets(workbook, kind: ReportKind):
    if kind == ReportKind.CHEMICAL_USAGE:
        return [sheet for sheet in workbook.worksheets if sheet["F2"].value == "填表日期："]
    return [sheet for sheet in workbook.worksheets if sheet["G4"].value == "日期："]


def _prepare_daily_sheets(workbook, kind: ReportKind, dates: list[dt.date]):
    sheets = _daily_sheets(workbook, kind)
    if not sheets:
        raise ValueError("模板中没有日期工作表")
    while len(sheets) < len(dates):
        sheets.append(workbook.copy_worksheet(sheets[-1]))
    for sheet in sheets[len(dates):]:
        workbook.remove(sheet)
    sheets = sheets[:len(dates)]
    old_titles = [sheet.title for sheet in sheets]
    for index, sheet in enumerate(sheets, start=1):
        sheet.title = f"_日期_{index}"
    for sheet, value in zip(sheets, dates):
        sheet.title = f"{value.month}.{value.day}"
        address = "G2" if kind == ReportKind.CHEMICAL_USAGE else "I4"
        sheet[address] = f"{value.year}.{value.month}.{value.day}"
    for index, sheet in enumerate(sheets):
        if index == 0:
            continue
        previous = sheets[index - 1].title
        for row in sheet.iter_rows():
            for cell in row:
                if isinstance(cell.value, str) and cell.value.startswith("=") and "!" in cell.value:
                    cell.value = _replace_sheet_references(cell.value, previous)
    return sheets


def _rename_in_column(sheets, column: int, old_name: str, new_name: str) -> None:
    found = False
    for sheet in sheets:
        for row in range(1, sheet.max_row + 1):
            cell = sheet.cell(row, column)
            if cell.value == old_name:
                cell.value = new_name
                found = True
    if not found:
        raise ValueError(f"没有找到供应商：{old_name}")


def _add_chemical_supplier(sheets, name: str) -> None:
    for index, sheet in enumerate(sheets):
        total_row = next((cell.row for cell in sheet["A"] if cell.value == "合计"), None)
        if total_row is None:
            raise ValueError("化肥矿石用量表缺少合计行")
        _insert_row_preserving_merges(sheet, total_row)
        _copy_row_style(sheet, total_row - 1, total_row, 7)
        for column in range(1, 8):
            sheet.cell(total_row, column).value = None
        sheet.cell(total_row, 2).value = name
        sheet.cell(total_row, 5).value = f"=D{total_row}" if index == 0 else f"=D{total_row}+'{sheets[index - 1].title}'!E{total_row}"
        sheet.cell(total_row, 6).value = f"=D{total_row}" if index == 0 else f"=D{total_row}+'{sheets[index - 1].title}'!F{total_row}"
        moved_total = total_row + 1
        for column, letter in ((4, "D"), (5, "E"), (6, "F")):
            sheet.cell(moved_total, column).value = f"=SUM({letter}4:{letter}{total_row})"
        category_merge = next(
            (merged for merged in list(sheet.merged_cells.ranges) if merged.min_col == 1 and merged.max_col == 1 and merged.min_row == 4),
            None,
        )
        if category_merge is not None:
            sheet.unmerge_cells(str(category_merge))
        sheet.merge_cells(start_row=4, start_column=1, end_row=total_row, end_column=1)
        sheet["A4"] = "矿石"


def _delete_chemical_supplier(sheets, name: str) -> None:
    for sheet in sheets:
        row = next((cell.row for cell in sheet["B"] if cell.value == name), None)
        if row is None:
            raise ValueError(f"没有找到供应商：{name}")
        _delete_row_preserving_merges(sheet, row)
        total_row = next(cell.row for cell in sheet["A"] if cell.value == "合计")
        for column, letter in ((4, "D"), (5, "E"), (6, "F")):
            sheet.cell(total_row, column).value = f"=SUM({letter}4:{letter}{total_row - 1})"
        sheet["A4"] = "矿石"


def _add_dispatch_supplier(sheets, name: str) -> None:
    for index, sheet in enumerate(sheets):
        total_row = next(
            (row for row in range(15, sheet.max_row + 1) if sheet.cell(row, 1).value is None and sheet.cell(row - 1, 1).value is not None),
            None,
        )
        if total_row is None:
            raise ValueError("调度生产信息表缺少外购汇总行")
        sheet.insert_rows(total_row)
        _copy_row_style(sheet, total_row - 1, total_row, sheet.max_column)
        for column in range(1, sheet.max_column + 1):
            sheet.cell(total_row, column).value = None
        sheet.cell(total_row, 1).value = name
        sheet.cell(total_row, 3).value = f"=B{total_row}" if index == 0 else f"=B{total_row}+'{sheets[index - 1].title}'!C{total_row}"
        sheet.cell(total_row, 4).value = f"=B{total_row}" if index == 0 else f"=B{total_row}+'{sheets[index - 1].title}'!D{total_row}"
        sheet.cell(total_row + 1, 2).value = f"=SUM(B14:B{total_row})"
        _extend_section_merge(sheet, 5, 14, total_row)


def _delete_dispatch_supplier(sheets, name: str) -> None:
    for sheet in sheets:
        row = next((cell.row for cell in sheet["A"] if cell.value == name), None)
        if row is None:
            raise ValueError(f"没有找到供应商：{name}")
        _delete_row_preserving_merges(sheet, row)
        total_row = next(
            row for row in range(15, sheet.max_row + 1)
            if sheet.cell(row, 1).value is None and sheet.cell(row - 1, 1).value is not None
        )
        sheet.cell(total_row, 2).value = f"=SUM(B14:B{total_row - 1})"


def _clear_chemical_inputs(sheets) -> None:
    for sheet in sheets:
        total_row = next((cell.row for cell in sheet["A"] if cell.value == "合计"), sheet.max_row + 1)
        for row in range(4, total_row):
            sheet.cell(row, 3).value = None
            sheet.cell(row, 4).value = None
            sheet.cell(row, 7).value = None


def _clear_dispatch_inputs(sheets) -> None:
    for sheet in sheets:
        for address in ("A5", "A7", "A9"):
            sheet[address] = None
        total_row = next(
            (row for row in range(15, sheet.max_row + 1) if sheet.cell(row, 1).value is None and sheet.cell(row - 1, 1).value is not None),
            sheet.max_row + 1,
        )
        sheet["E14"] = None
        for row in range(14, total_row):
            sheet.cell(row, 2).value = None
        for row in range(14, min(total_row, 26)):
            sheet.cell(row, 8).value = None


def _clear_inventory_inputs(sheet) -> None:
    note_row = next((cell.row for cell in sheet["B"] if cell.value == "填表说明："), sheet.max_row + 1)
    for row in range(3, note_row):
        for column in range(6, 37):
            cell = sheet.cell(row, column)
            if cell.data_type != "f":
                cell.value = None


def _active_inventory_sheet(workbook):
    candidates = []
    for sheet in workbook.worksheets:
        dates = [sheet.cell(2, column).value for column in range(6, 37)]
        actual = [value.date() if isinstance(value, dt.datetime) else value for value in dates if isinstance(value, (dt.date, dt.datetime))]
        if actual:
            candidates.append((max(actual), sheet))
    if not candidates:
        raise ValueError("产消存计划表中没有日期区域")
    return max(candidates, key=lambda item: item[0])[1]


def _update_inventory(workbook, dates: list[dt.date], changes: list[SupplierChange]) -> None:
    sheet = _active_inventory_sheet(workbook)
    for offset, column in enumerate(range(6, 37)):
        sheet.cell(2, column).value = dates[offset] if offset < len(dates) else None
    for change in changes:
        if change.action == "rename":
            boundaries = (3, 15) if change.area == "生产日累" else (50, 56)
            found = False
            for row in range(boundaries[0], boundaries[1] + 1):
                if sheet.cell(row, 2).value == change.old_name:
                    sheet.cell(row, 2).value = change.new_name
                    found = True
            if not found:
                raise ValueError(f"没有在{change.area}区域找到供应商：{change.old_name}")
        elif change.action == "delete":
            if change.area == "生产日累":
                end = next(cell.row for cell in sheet["A"] if cell.value == "库存量") - 1
                boundaries = (3, end)
            else:
                end = next(cell.row for cell in sheet["B"] if cell.value == "填表说明：") - 1
                start = next(cell.row for cell in sheet["A"] if cell.value == "当日配送量(实发量)")
                boundaries = (start, end)
            row = next(
                (row for row in range(boundaries[0], min(boundaries[1], sheet.max_row) + 1) if sheet.cell(row, 2).value == change.old_name),
                None,
            )
            if row is None:
                raise ValueError(f"没有在{change.area}区域找到供应商：{change.old_name}")
            _delete_row_preserving_merges(sheet, row)
        elif change.area == "生产日累":
            section_row = next(cell.row for cell in sheet["A"] if cell.value == "生产日累")
            inventory_row = next(cell.row for cell in sheet["A"] if cell.value == "库存量")
            _insert_row_preserving_merges(sheet, inventory_row)
            _copy_row_style(sheet, inventory_row - 1, inventory_row, 38)
            for column in range(1, 39):
                sheet.cell(inventory_row, column).value = None
            sheet.cell(inventory_row, 2).value = change.new_name
            sheet.cell(inventory_row, 37).value = f"=SUM(F{inventory_row}:AJ{inventory_row})"
            _merge_inventory_supplier_row(sheet, inventory_row, section_row)
        else:
            section_row = next(cell.row for cell in sheet["A"] if cell.value == "当日配送量(实发量)")
            note_row = next(cell.row for cell in sheet["B"] if cell.value == "填表说明：")
            target_row = note_row - 1
            if sheet.cell(target_row, 2).value:
                _insert_row_preserving_merges(sheet, note_row)
                target_row = note_row
            _copy_row_style(sheet, target_row - 1, target_row, 38)
            for column in range(1, 39):
                sheet.cell(target_row, column).value = None
            sheet.cell(target_row, 2).value = change.new_name
            sheet.cell(target_row, 37).value = f"=SUM(F{target_row}:AJ{target_row})"
            _merge_inventory_supplier_row(sheet, target_row, section_row)
            note_merge = next(
                (
                    merged for merged in list(sheet.merged_cells.ranges)
                    if merged.min_col == 38 == merged.max_col and merged.max_row == target_row - 1
                ),
                None,
            )
            if note_merge is not None:
                _extend_section_merge(sheet, 38, note_merge.min_row, target_row)
    _ensure_inventory_monthly_formulas(sheet)


def update_legacy_report(
    source: str | Path,
    output: str | Path,
    start: dt.date,
    end: dt.date,
    supplier_changes: list[SupplierChange] | None = None,
    title: str | None = None,
    row_rules: list[DimensionRule] | None = None,
    column_rules: list[DimensionRule] | None = None,
) -> ReportKind:
    values = _dates(start, end)
    changes = supplier_changes or []
    kind = detect_report(source)
    workbook = load_workbook(source, data_only=False)
    if kind == ReportKind.INVENTORY_PLAN:
        _update_inventory(workbook, values, changes)
        _clear_inventory_inputs(_active_inventory_sheet(workbook))
        if title:
            _active_inventory_sheet(workbook)["A1"] = title
    else:
        sheets = _prepare_daily_sheets(workbook, kind, values)
        if title:
            address = "A1" if kind == ReportKind.CHEMICAL_USAGE else "A3"
            for sheet in sheets:
                sheet[address] = title
        for change in changes:
            column = 2 if kind == ReportKind.CHEMICAL_USAGE else 1
            if change.action == "rename":
                _rename_in_column(sheets, column, change.old_name, change.new_name)
            elif change.action == "delete" and kind == ReportKind.CHEMICAL_USAGE:
                _delete_chemical_supplier(sheets, change.old_name)
            elif change.action == "delete":
                _delete_dispatch_supplier(sheets, change.old_name)
            elif kind == ReportKind.CHEMICAL_USAGE:
                _add_chemical_supplier(sheets, change.new_name)
            else:
                _add_dispatch_supplier(sheets, change.new_name)
        if kind == ReportKind.CHEMICAL_USAGE:
            _clear_chemical_inputs(sheets)
        else:
            _clear_dispatch_inputs(sheets)
    target_sheets = [_active_inventory_sheet(workbook)] if kind == ReportKind.INVENTORY_PLAN else _daily_sheets(workbook, kind)
    for sheet in target_sheets:
        for rule in row_rules or []:
            height = row_height(rule.value, rule.unit)
            for number in parse_rows(rule.selection):
                sheet.row_dimensions[number].height = height
        for rule in column_rules or []:
            width = column_width(rule.value, rule.unit)
            for name in parse_columns(rule.selection):
                sheet.column_dimensions[name].width = width
    workbook.calculation.fullCalcOnLoad = True
    workbook.calculation.forceFullCalc = True
    workbook.calculation.calcMode = "auto"
    output_path = Path(output).expanduser().resolve()
    output_path.parent.mkdir(parents=True, exist_ok=True)
    monthly_cache_cells = []
    inventory_sheet = None
    if kind == ReportKind.INVENTORY_PLAN:
        inventory_sheet = _active_inventory_sheet(workbook)
        monthly_cache_cells = [
            cell.coordinate for cell in inventory_sheet["AK"]
            if isinstance(cell.value, str) and cell.value.startswith("=")
        ]
    workbook.save(output_path)
    if monthly_cache_cells and inventory_sheet is not None:
        _cache_formula_zeros(output_path, inventory_sheet.path, monthly_cache_cells)
    return kind
